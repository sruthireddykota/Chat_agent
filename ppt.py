# Create a PowerPoint presentation for the Multi-Agent Chatbot System
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Pt

# Initialize presentation
prs = Presentation()

# Helper function to add slide
def add_bullet_slide(title, content_lines):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    content = slide.placeholders[1]
    content.text = content_lines[0]
    for line in content_lines[1:]:
        p = content.text_frame.add_paragraph()
        p.text = line
        p.level = 1

# Slide 1: Title Slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Multi-Agent Chatbot System"
subtitle = slide.placeholders[1]
subtitle.text = "AI-Powered Modular Chat Platform\nPresented by Sruthi Kota"

# Slide 2: Problem Statement
add_bullet_slide(
    "Problem Statement",
    [
        "Need for an intelligent, scalable AI assistant system",
        "Single-agent systems lack specialization",
        "Limited memory, retrieval accuracy, and real-time research capability"
    ]
)

# Slide 3: Solution Overview
add_bullet_slide(
    "Solution Overview",
    [
        "Multi-Agent AI Chat System",
        "Specialized agents for different tasks",
        "Integrated with Azure LLM, Redis, MongoDB, Qdrant, and MCP Server"
    ]
)

# Slide 4: System Architecture
add_bullet_slide(
    "System Architecture",
    [
        "Streamlit Frontend",
        "Azure OpenAI for LLM processing",
        "Redis for short-term memory",
        "MongoDB for persistence",
        "Qdrant for vector search",
        "MCP Server for tool orchestration"
    ]
)

# Slide 5: Dashboard
slide_layout = prs.slide_layouts[5]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Application Dashboard"
# slide.shapes.add_picture("/documentation/Demo Images/Dashboard.pdf", Inches(1), Inches(1.5), width=Inches(6))

# Slide 6: Chat Interface
slide_layout = prs.slide_layouts[5]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Chat Interface"
# slide.shapes.add_picture("/documentation/Demo Images/Chatbot.pdf", Inches(1), Inches(1.5), width=Inches(6))

# Slide 7: Agents Overview
add_bullet_slide(
    "Agents Overview",
    [
        "Generic Agent – General conversation & file analysis",
        "Coder Agent – Code debugging & execution",
        "RAG Agent – Retrieval-based document Q&A",
        "Researcher Agent – Real-time web research"
    ]
)

# Slide 8: Document Parser Pipeline
slide_layout = prs.slide_layouts[5]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Document Parser Pipeline"
# slide.shapes.add_picture("/documentation/Demo Images/Document Parser.pdf", Inches(1), Inches(1.5), width=Inches(6))

# Slide 9: Knowledge Management
slide_layout = prs.slide_layouts[5]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Knowledge Management"
slide.shapes.add_picture(r"documentation/Demo Images/KM dashboard.png", Inches(1), Inches(1.5), width=Inches(6))

# Slide 10: Key Features
add_bullet_slide(
    "Key Features",
    [
        "Multi-agent modular architecture",
        "Real-time streaming responses",
        "Retrieval-augmented accuracy",
        "Code execution support",
        "Persistent session management",
        "Docker-based scalable deployment"
    ]
)

# Slide 11: Conclusion
add_bullet_slide(
    "Conclusion",
    [
        "Combines multiple AI capabilities in one platform",
        "Scalable, modular, and production-ready",
        "Suitable for enterprise AI assistant systems"
    ]
)

# Save presentation
file_path =r"documentation/Demo Images/Multi_Agent_Chatbot_Presentation.pptx"
prs.save(file_path)

file_path
